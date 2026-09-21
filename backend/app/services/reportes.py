from datetime import datetime
from io import BytesIO
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as RLImage,
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

# Paleta Pixel Store (sin neón): azul oscuro casi negro, azul verdoso, blanco
AZUL_OSCURO = colors.HexColor("#0f172a")
AZUL_MEDIO = colors.HexColor("#0e7490")
GRIS_TEXTO = colors.HexColor("#475569")
GRIS_SUAVE = colors.HexColor("#f1f5f8")
BLANCO = colors.white
NEGRO = colors.HexColor("#0f172a")

AZUL_OSCURO_HEX = "0F172A"
AZUL_MEDIO_HEX = "0E7490"
GRIS_SUAVE_HEX = "F1F5F8"
BLANCO_HEX = "FFFFFF"

HEADERS = ["Venta", "Cliente", "Fecha", "Subtotal", "Impuestos", "Total", "Estado"]

STATIC_DIR = Path(__file__).resolve().parent.parent / "static"
LOGO_RUTA = STATIC_DIR / "logoPixel_circle.png"
if not LOGO_RUTA.exists():
    LOGO_RUTA = STATIC_DIR / "logoPixel.png"

TITULO = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=22, textColor=AZUL_MEDIO, leading=26)
SUBTITULO = ParagraphStyle("subtitulo", fontName="Helvetica", fontSize=10, textColor=GRIS_TEXTO, leading=14)
CELDA = ParagraphStyle("celda", fontName="Helvetica", fontSize=9, textColor=NEGRO, leading=12)
CELDA_BOLD = ParagraphStyle("celda_bold", fontName="Helvetica-Bold", fontSize=9, textColor=NEGRO, leading=12)
CELDA_TITULO = ParagraphStyle("celda_titulo", fontName="Helvetica-Bold", fontSize=9, textColor=BLANCO, leading=12)
CELDA_TOTAL = ParagraphStyle("celda_total", fontName="Helvetica-Bold", fontSize=9, textColor=BLANCO, leading=12)


def _value(row, name, default=""):
    if isinstance(row, dict):
        return row.get(name, default)
    return getattr(row, name, default)


def _rows(rows):
    result = []
    for row in rows:
        result.append({
            "venta": _value(row, "venta_id", _value(row, "id", "")),
            "cliente": _value(row, "cliente_id", ""),
            "fecha": _value(row, "fecha", ""),
            "subtotal": Decimal(str(_value(row, "subtotal", 0) or 0)),
            "impuestos": Decimal(str(_value(row, "impuestos", 0) or 0)),
            "total": Decimal(str(_value(row, "total", 0) or 0)),
            "estado": _value(row, "estado", ""),
        })
    return result


def _logo_pdf() -> RLImage | None:
    if not LOGO_RUTA.exists():
        return None
    try:
        from PIL import Image as PILImage
        imagen = PILImage.open(str(LOGO_RUTA))
        if imagen.mode != "RGBA":
            imagen = imagen.convert("RGBA")
        buffer = BytesIO()
        imagen.save(buffer, format="PNG")
        buffer.seek(0)
        rl = RLImage(buffer, width=0.6 * inch, height=0.6 * inch)
        rl.hAlign = "LEFT"
        return rl
    except Exception:
        return None


def _grafica_lineas_pdf(values):
    """Gráfica de líneas de ganancias (totales) con colores de la página."""
    from reportlab.graphics import renderPDF
    from reportlab.graphics.shapes import Drawing, Line, PolyLine, Rect, String

    if not values:
        return None
    ancho, alto = 6.9 * inch, 2.2 * inch
    d = Drawing(ancho, alto)
    margen_x, margen_y = 34, 26
    plot_w = ancho - margen_x * 2
    plot_h = alto - margen_y * 2

    maximos = max(float(v["total"]) for v in values) or 1
    puntos = []
    n = len(values)
    for i, v in enumerate(values):
        x = margen_x + (i * plot_w) / max(n - 1, 1)
        y = margen_y + (float(v["total"]) / maximos) * plot_h
        puntos.append((x, y))

    # fondo y ejes (azul oscuro casi negro + gris)
    d.add(Rect(margen_x, margen_y, plot_w, plot_h, fillColor=colors.HexColor("#f8fafc"), strokeColor=AZUL_OSCURO, strokeWidth=0.8))
    d.add(Line(margen_x, margen_y, margen_x + plot_w, margen_y, strokeColor=AZUL_OSCURO, strokeWidth=1))
    d.add(Line(margen_x, margen_y, margen_x, margen_y + plot_h, strokeColor=AZUL_OSCURO, strokeWidth=1))

    # líneas de referencia horizontales
    for f in (0.25, 0.5, 0.75, 1.0):
        y = margen_y + f * plot_h
        d.add(Line(margen_x, y, margen_x + plot_w, y, strokeColor=colors.HexColor("#cbd5e1"), strokeDashArray=[2, 3]))

    # área bajo la curva (azul verdoso tenue)
    area = [puntos[0][0], margen_y] + [coord for p in puntos for coord in p] + [puntos[-1][0], margen_y]
    from reportlab.graphics.shapes import Polygon
    d.add(Polygon(area, fillColor=AZUL_MEDIO, strokeColor=None, strokeWidth=0, fillOpacity=0.18))

    # línea principal (azul verdoso)
    d.add(PolyLine(puntos, strokeColor=AZUL_MEDIO, strokeWidth=2.2, strokeLineJoin=1, strokeLineCap=1))

    # puntos y etiquetas
    for i, (x, y) in enumerate(puntos):
        d.add(RL_Circle(x, y, 2.6, fillColor=AZUL_OSCURO, strokeColor=BLANCO, strokeWidth=1))
        etiqueta = values[i]["fecha"]
        d.add(String(x, margen_y - 12, str(etiqueta)[:10], fontSize=7, textAnchor="middle", fillColor=GRIS_TEXTO))
        d.add(String(x, y + 5, f"${float(values[i]['total']):,.0f}", fontSize=6.5, textAnchor="middle", fillColor=AZUL_OSCURO))

    return d


# helper para círculos en el drawing (reportlab no exporta Circle directamente en shapes top-level siempre)
try:
    from reportlab.graphics.shapes import Circle as RL_Circle
except Exception:
    def RL_Circle(x, y, r, fillColor=None, strokeColor=None, strokeWidth=0):
        from reportlab.graphics.shapes import Circle
        return Circle(x, y, r, fillColor=fillColor, strokeColor=strokeColor, strokeWidth=strokeWidth)


def generar_pdf_ventas(rows):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=36, leftMargin=36, topMargin=30, bottomMargin=30)
    values = _rows(rows)

    # Cabecera con logo + nombre
    logo = _logo_pdf()
    if logo is not None:
        celda_izq = [logo, Paragraph("PIXEL STORE", TITULO), Paragraph("Tecnología que conecta contigo", SUBTITULO)]
    else:
        celda_izq = [Paragraph("PIXEL STORE", TITULO), Paragraph("Tecnología que conecta contigo", SUBTITULO)]
    cabecera = Table([[celda_izq, [Paragraph("REPORTE DE VENTAS", ParagraphStyle("rep", fontName="Helvetica-Bold", fontSize=14, textColor=BLANCO, leading=18)), Paragraph(f"Generado: {datetime.now():%Y-%m-%d %H:%M}", ParagraphStyle("rep_sub", fontName="Helvetica", fontSize=9, textColor=colors.HexColor("#cbd5e1"), leading=12))]]], colWidths=[4.4 * inch, 3.6 * inch])
    cabecera.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (1, 0), (1, 0), AZUL_OSCURO),
        ("TOPPADDING", (1, 0), (1, 0), 10),
        ("BOTTOMPADDING", (1, 0), (1, 0), 10),
        ("LEFTPADDING", (1, 0), (1, 0), 16),
        ("ROUNDEDCORNERS", [8, 8, 8, 8]),
    ]))

    # Tabla de datos
    data = [[Paragraph(h, CELDA_TITULO) for h in HEADERS]]
    for v in values:
        data.append([
            Paragraph(str(v["venta"]), CELDA),
            Paragraph(str(v["cliente"]), CELDA),
            Paragraph(str(v["fecha"])[:19], CELDA),
            Paragraph(f"${v['subtotal']:,.0f}", CELDA),
            Paragraph(f"${v['impuestos']:,.0f}", CELDA),
            Paragraph(f"${v['total']:,.0f}", CELDA),
            Paragraph(str(v["estado"]), CELDA),
        ])
    total = sum(v["total"] for v in values)
    data.append([
        Paragraph("", CELDA_TOTAL), Paragraph("", CELDA_TOTAL), Paragraph("TOTAL", CELDA_TOTAL),
        Paragraph(f"${sum(v['subtotal'] for v in values):,.0f}", CELDA_TOTAL),
        Paragraph(f"${sum(v['impuestos'] for v in values):,.0f}", CELDA_TOTAL),
        Paragraph(f"${total:,.0f}", CELDA_TOTAL),
        Paragraph("", CELDA_TOTAL),
    ])
    tabla = Table(data, repeatRows=1, colWidths=[0.9 * inch, 1.0 * inch, 1.8 * inch, 1.3 * inch, 1.3 * inch, 1.3 * inch, 1.1 * inch])
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), AZUL_OSCURO),
        ("TEXTCOLOR", (0, 0), (-1, 0), BLANCO),
        ("ALIGN", (3, 1), (5, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [BLANCO, GRIS_SUAVE]),
        ("BACKGROUND", (0, -1), (-1, -1), AZUL_MEDIO),
        ("TEXTCOLOR", (0, -1), (-1, -1), BLANCO),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))

    story = [
        cabecera,
        Spacer(1, 16),
        tabla,
        Spacer(1, 20),
        Paragraph("Gráfica de ganancias", ParagraphStyle("graf", fontName="Helvetica-Bold", fontSize=13, textColor=AZUL_OSCURO, leading=16)),
        Spacer(1, 8),
    ]
    grafica = _grafica_lineas_pdf(values)
    if grafica is not None:
        story.append(grafica)

    doc.build(story)
    buffer.seek(0)
    return buffer


def generar_xlsx_ventas(rows):
    buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"

    # Logo + nombre en cabecera
    try:
        if LOGO_RUTA.exists():
            from openpyxl.drawing.image import Image as XLImage
            logo = XLImage(str(LOGO_RUTA))
            logo.width = 48
            logo.height = 48
            sheet.add_image(logo, "A1")
    except Exception:
        pass

    sheet["B1"] = "PIXEL STORE"
    sheet["B1"].font = Font(bold=True, size=20, color=AZUL_MEDIO_HEX)
    sheet["B1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 52

    sheet["B2"] = f"Reporte de ventas · Generado: {datetime.now():%Y-%m-%d %H:%M}"
    sheet["B2"].font = Font(size=10, color="475569")

    fila_encabezado = 4
    for col, titulo in enumerate(HEADERS, start=1):
        celda = sheet.cell(fila_encabezado, col, titulo)
        celda.font = Font(bold=True, color=BLANCO_HEX)
        celda.fill = PatternFill("solid", fgColor=AZUL_OSCURO_HEX)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    fila_encabezado += 1

    thin = Side(style="thin", color="E2E8F0")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    fila_inicio = fila_encabezado
    for v in _rows(rows):
        valores = [v["venta"], v["cliente"], str(v["fecha"])[:19], float(v["subtotal"]), float(v["impuestos"]), float(v["total"]), v["estado"]]
        for col, valor in enumerate(valores, start=1):
            celda = sheet.cell(fila_encabezado, col, valor)
            celda.font = Font(size=10, color="0F172A")
            celda.border = borde
            if 4 <= col <= 6:
                celda.alignment = Alignment(horizontal="right")
        fila_encabezado += 1

    # Fila total
    total_fila = fila_encabezado
    sheet.cell(total_fila, 3, "TOTAL").font = Font(bold=True, color=BLANCO_HEX)
    for col in range(1, 8):
        celda_t = sheet.cell(total_fila, col)
        celda_t.fill = PatternFill("solid", fgColor=AZUL_MEDIO_HEX)
        celda_t.font = Font(bold=True, color=BLANCO_HEX)
    for col in range(4, 7):
        col_letra = chr(64 + col)
        sheet.cell(total_fila, col, f"=SUM({col_letra}{fila_inicio}:{col_letra}{total_fila - 1})")
        sheet.cell(total_fila, col).font = Font(bold=True, color=BLANCO_HEX)
        sheet.cell(total_fila, col).alignment = Alignment(horizontal="right")

    for fila in range(fila_inicio, total_fila + 1):
        for col in range(4, 7):
            sheet.cell(fila, col).number_format = '$#,##0.00'

    # Gráfica de líneas de ganancias (azul verdoso)
    chart = LineChart()
    chart.title = "Gráfica de ganancias"
    chart.style = 12
    chart.height = 9
    chart.width = 22
    chart.y_axis.title = "Total"
    chart.x_axis.title = "Venta"
    datos_ref = Reference(sheet, min_col=6, min_row=fila_encabezado - 1, max_row=total_fila - 1)  # col F (total)
    categorias_ref = Reference(sheet, min_col=1, min_row=fila_encabezado, max_row=total_fila - 1)  # venta
    chart.add_data(datos_ref, titles_from_data=True)
    chart.set_categories(categorias_ref)
    try:
        serie = chart.series[0]
        serie.graphicalProperties.line.solidFill = AZUL_MEDIO_HEX
        serie.graphicalProperties.line.width = 22000
        serie.marker.symbol = "circle"
        serie.marker.graphicalProperties.solidFill = AZUL_OSCURO_HEX
        serie.smooth = False
    except Exception:
        pass
    sheet.add_chart(chart, f"A{total_fila + 2}")

    for col, ancho in {"A": 12, "B": 14, "C": 22, "D": 15, "E": 15, "F": 15, "G": 15}.items():
        sheet.column_dimensions[col].width = ancho

    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def filas_ventas(ventas):
    return ventas
