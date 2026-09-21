from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from app.core.dependencies import obtener_usuario_actual
from app.database.database import get_db
from app.models import DetalleFactura, Factura, Venta
from app.schemas.factura import FacturaResponse
router = APIRouter(prefix="/api/facturas", tags=["Facturas"])
def _crear_desde_venta(venta, db):
    if venta.factura: return venta.factura
    factura = Factura(venta_id=venta.id, numero=f"PS-{datetime.utcnow():%Y%m%d}-{venta.id:06d}", subtotal=venta.subtotal, impuestos=venta.impuestos, total=venta.total, estado="emitida")
    factura.detalles = [DetalleFactura(producto_id=d.producto_id, servicio_id=d.servicio_id, descripcion=d.nombre_item, cantidad=d.cantidad, precio_unitario=d.precio_unitario, subtotal=d.subtotal) for d in venta.detalles]
    db.add(factura); db.commit(); db.refresh(factura); return factura
@router.post("/venta/{venta_id}", response_model=FacturaResponse, status_code=201)
def crear_factura(venta_id: int, db: Session = Depends(get_db), actual=Depends(obtener_usuario_actual)):
    venta = db.query(Venta).options(joinedload(Venta.detalles), joinedload(Venta.factura)).filter(Venta.id == venta_id).first()
    if not venta: raise HTTPException(404, "Venta no encontrada")
    if actual.rol.nombre == "cliente" and venta.cliente_id != actual.id: raise HTTPException(403, "No tienes permiso")
    return _crear_desde_venta(venta, db)
def _serializar_factura(factura):
    """Añade operador (quién la registró) y tipo (venta/pedido) a la respuesta."""
    venta = factura.venta
    pedido = factura.pedido
    operador = None
    tipo = None
    if venta and venta.operador:
        operador = {
            "id": venta.operador.id,
            "nombre": venta.operador.nombre,
            "apellido": venta.operador.apellido,
        }
        tipo = "venta"
    elif pedido:
        tipo = "pedido"
    return {
        "id": factura.id,
        "venta_id": factura.venta_id,
        "pedido_id": factura.pedido_id,
        "numero": factura.numero,
        "fecha": factura.fecha,
        "subtotal": factura.subtotal,
        "impuestos": factura.impuestos,
        "total": factura.total,
        "estado": factura.estado,
        "detalles": factura.detalles,
        "operador": operador,
        "tipo": tipo,
    }


@router.get("", response_model=list[FacturaResponse])
def listar_facturas(
    numero: str | None = Query(default=None),
    estado: str | None = Query(default=None),
    fecha: str | None = Query(default=None),
    cliente_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    actual=Depends(obtener_usuario_actual),
):
    from datetime import date, timedelta
    from app.models import Pedido
    q = db.query(Factura).outerjoin(Venta).outerjoin(Pedido, Pedido.id == Factura.pedido_id).options(joinedload(Factura.detalles), joinedload(Factura.venta).joinedload(Venta.operador)).order_by(Factura.fecha.desc())
    if actual.rol.nombre == "cliente":
        q = q.filter((Venta.cliente_id == actual.id) | (Pedido.usuario_id == actual.id))
    elif cliente_id:
        q = q.filter((Venta.cliente_id == cliente_id) | (Pedido.usuario_id == cliente_id))
    if numero:
        q = q.filter(Factura.numero.contains(numero))
    if estado:
        q = q.filter(Factura.estado == estado)
    if fecha:
        try:
            dia = date.fromisoformat(fecha)
            q = q.filter(Factura.fecha >= datetime.combine(dia, datetime.min.time()), Factura.fecha < datetime.combine(dia + timedelta(days=1), datetime.min.time()))
        except ValueError:
            pass
    facturas = q.all()
    return [_serializar_factura(f) for f in facturas]
def _obtener_factura_orm(factura_id: int, db: Session, actual):
    from app.models import Pedido
    factura = db.query(Factura).outerjoin(Venta).outerjoin(Pedido, Pedido.id == Factura.pedido_id).options(joinedload(Factura.detalles), joinedload(Factura.venta).joinedload(Venta.operador), joinedload(Factura.pedido)).filter(Factura.id == factura_id).first()
    if not factura:
        raise HTTPException(404, "Factura no encontrada")
    if actual.rol.nombre == "cliente":
        pertenece = (factura.venta and factura.venta.cliente_id == actual.id) or (factura.pedido and factura.pedido.usuario_id == actual.id)
        if not pertenece:
            raise HTTPException(403, "No tienes permiso")
    return factura


@router.get("/{factura_id}", response_model=FacturaResponse)
def obtener_factura(factura_id: int, db: Session = Depends(get_db), actual=Depends(obtener_usuario_actual)):
    factura = _obtener_factura_orm(factura_id, db, actual)
    return _serializar_factura(factura)

@router.get("/{factura_id}/pdf")
def descargar_factura(factura_id: int, db: Session = Depends(get_db), actual=Depends(obtener_usuario_actual)):
    from app.services.facturacion_pdf import factura_pdf
    factura_orm = _obtener_factura_orm(factura_id, db, actual)
    return StreamingResponse(factura_pdf(factura_orm), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{factura_orm.numero}.pdf"'})


def _factura_excel_bytes(factura):
    from io import BytesIO
    from pathlib import Path

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    AZUL_OSCURO = "0F172A"
    CIAN = "22D3EE"
    AZUL_MEDIO = "0E7490"
    GRIS_SUAVE = "F1F5F8"
    BLANCO = "FFFFFF"

    LOGO = Path(__file__).resolve().parent.parent / "static" / "logoPixel.png"

    def _porcentaje_impuesto():
        subtotal = float(factura.subtotal or 0)
        impuestos = float(factura.impuestos or 0)
        if subtotal <= 0 or impuestos <= 0:
            return 0
        return int(round((impuestos / subtotal) * 100))

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Factura"

    # Logo circular pequeño + nombre al lado en la cabecera
    LOGO_CIRCLE = Path(__file__).resolve().parent.parent / "static" / "logoPixel_circle.png"
    ruta_logo = LOGO_CIRCLE if LOGO_CIRCLE.exists() else LOGO
    try:
        if ruta_logo.exists():
            logo = XLImage(str(ruta_logo))
            logo.width = 54
            logo.height = 54
            ws.add_image(logo, "A1")
    except Exception:
        pass

    ws["B1"] = "PIXEL STORE"
    ws["B1"].font = Font(bold=True, size=22, color=AZUL_MEDIO)
    ws["B1"].alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 60

    ws["B2"] = "Tecnología que conecta contigo"
    ws["B2"].font = Font(size=10, color="475569")
    ws["B2"].alignment = Alignment(horizontal="left")

    ws.merge_cells("A3:E3")
    ws["A3"] = ""
    ws.row_dimensions[3].height = 8

    ws.merge_cells("A4:E4")
    ws["A4"] = f"FACTURA  NO° {factura.numero}"
    ws["A4"].font = Font(bold=True, size=18, color=BLANCO)
    ws["A4"].fill = PatternFill("solid", fgColor=AZUL_OSCURO)
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 34

    venta = factura.venta
    pedido = factura.pedido
    cliente = (venta.cliente if venta else None) or (pedido.usuario if pedido else None)
    cliente_nombre = f"{cliente.nombre} {cliente.apellido}" if cliente else "Cliente"
    cliente_correo = cliente.correo if cliente else ""
    cliente_direccion = cliente.direccion if cliente else ""
    cliente_documento = f"{cliente.tipo_documento} {cliente.numero_documento}" if cliente else ""

    fila = 6
    ws.cell(fila, 1, "CLIENTE").font = Font(bold=True, size=9, color="475569")
    ws.cell(fila, 1, "CLIENTE").fill = PatternFill("solid", fgColor=GRIS_SUAVE)
    ws.cell(fila, 2, "Correo").font = Font(bold=True, size=9, color="475569")
    ws.cell(fila, 2, "Correo").fill = PatternFill("solid", fgColor=GRIS_SUAVE)
    fila += 1
    ws.cell(fila, 1, cliente_nombre).font = Font(size=11, color="0F172A")
    ws.cell(fila, 2, cliente_correo).font = Font(size=11, color="0F172A")
    fila += 1
    ws.cell(fila, 1, f"Dirección: {cliente_direccion}").font = Font(size=10, color="0F172A")
    ws.cell(fila, 2, f"Documento: {cliente_documento}").font = Font(size=10, color="0F172A")
    fila += 1
    operador = venta.operador if (venta and venta.operador) else None
    if operador:
        operador_nombre = f"{operador.nombre} {operador.apellido}"
        ws.cell(fila, 1, f"Añadida por: {operador_nombre}").font = Font(size=10, color="0F172A")
    fila += 1

    encabezados = ["Artículo", "Cantidad/Horas", "Precio/u", "Total"]
    for col, titulo in enumerate(encabezados, start=1):
        celda = ws.cell(fila, col, titulo)
        celda.font = Font(bold=True, color=BLANCO)
        celda.fill = PatternFill("solid", fgColor=AZUL_OSCURO)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    fila += 1

    thin = Side(style="thin", color="E2E8F0")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    for detalle in factura.detalles:
        valores = [
            detalle.descripcion,
            detalle.cantidad,
            float(detalle.precio_unitario),
            float(detalle.subtotal),
        ]
        for col, valor in enumerate(valores, start=1):
            celda = ws.cell(fila, col, valor)
            celda.font = Font(size=10, color="0F172A")
            celda.border = borde
            if col >= 2:
                celda.alignment = Alignment(horizontal="right")
        fila += 1

    fila += 1
    porcentaje = _porcentaje_impuesto()
    totales = [
        ("Subtotal:", float(factura.subtotal)),
        (f"Impuestos ({porcentaje} %):", float(factura.impuestos)),
        ("Precio total:", float(factura.total)),
    ]
    for etiqueta, valor in totales:
        ws.cell(fila, 3, etiqueta).font = Font(bold=True, size=11, color="0F172A")
        ws.cell(fila, 3).alignment = Alignment(horizontal="right")
        celda = ws.cell(fila, 4, valor)
        celda.font = Font(bold=True, size=11, color=AZUL_MEDIO)
        celda.alignment = Alignment(horizontal="right")
        fila += 1

    fila += 1
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=5)
    ws.cell(fila, 1, "¡GRACIAS!  ·  Visita pixelstore.com  ·  contacto@pixelstore.com").font = Font(bold=True, size=12, color=AZUL_MEDIO)
    ws.cell(fila, 1).alignment = Alignment(horizontal="center")
    ws.row_dimensions[fila].height = 28

    for col, ancho in {"A": 34, "B": 16, "C": 16, "D": 16, "E": 16}.items():
        ws.column_dimensions[col].width = ancho

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


@router.get("/{factura_id}/excel")
def descargar_factura_excel(factura_id: int, db: Session = Depends(get_db), actual=Depends(obtener_usuario_actual)):
    factura = _obtener_factura_orm(factura_id, db, actual)
    buffer = _factura_excel_bytes(factura)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{factura.numero}.xlsx"'},
    )
